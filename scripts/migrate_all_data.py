import openpyxl
import os
import sys
import json
import re
from datetime import datetime, date, timedelta

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\ASUS\Downloads\Patent Dashboard"
OUTPUT_DIR = os.path.join(BASE_DIR, "data")
os.makedirs(OUTPUT_DIR, exist_ok=True)

PLACEHOLDERS = {'', 'na', 'n/a', '-', '--', 'nil', 'none', 'not filed', 'not yet postdated', 'ps not yet filed'}
RULE_CUTOFF = date(2024, 3, 15)

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    return '' if s.lower() in PLACEHOLDERS or s.startswith('#') else s

def parse_date(val):
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    if s.lower() in PLACEHOLDERS or s.startswith('#'):
        return None
    m = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{4})', s)
    if m:
        try:
            d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            return date(y, mth, d).strftime('%Y-%m-%d')
        except ValueError:
            pass
    m2 = re.match(r'^(\d{4})[/-](\d{1,2})[/-](\d{1,2})', s)
    if m2:
        try:
            y, mth, d = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
            return date(y, mth, d).strftime('%Y-%m-%d')
        except ValueError:
            pass
    try:
        num = float(s)
        if 20000 < num < 60000:
            dt = datetime.fromordinal(datetime(1900, 1, 1).toordinal() + int(num) - 2)
            return dt.strftime('%Y-%m-%d')
    except ValueError:
        pass
    return None

def add_months_date(d_str, months):
    if not d_str:
        return None
    try:
        dt = datetime.strptime(d_str, '%Y-%m-%d').date()
        m = dt.month - 1 + months
        year = dt.year + m // 12
        month = m % 12 + 1
        day = min(dt.day, [31, 29 if year%4==0 and (year%100!=0 or year%400==0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month-1])
        return date(year, month, day).strftime('%Y-%m-%d')
    except Exception:
        return None

def add_days_date(d_str, days):
    if not d_str:
        return None
    try:
        dt = datetime.strptime(d_str, '%Y-%m-%d').date()
        return (dt + timedelta(days=days)).strftime('%Y-%m-%d')
    except Exception:
        return None

MAHE_PREFIXES = {
    'AR', 'BR', 'CD', 'CH', 'DA', 'DH', 'HK', 'HP', 'KM', 'MA', 'MB', 'MCODS',
    'MK', 'MN', 'MP', 'MS', 'MT', 'PH', 'SP', 'AP', 'LT', 'MI', 'DST'
}

def classify_prefix(code):
    m = re.match(r'^[A-Za-z_]+', str(code).strip())
    prefix = m.group(0).upper() if m else ''
    is_mahe = prefix in MAHE_PREFIXES
    category = 'INTERNAL_FILING'
    if prefix.startswith('EX'):
        category = 'EXTERNAL_TAKEOVER'
    elif prefix in {'DA', 'KM', 'MB', 'MS', 'MI'}:
        category = 'WORKSHOP_IDF'
    return prefix, is_mahe, category

def split_numbered_list(text):
    if not text:
        return []
    s = str(text).replace('\r\n', '\n').strip()
    pieces = re.split(r'\b\d+\s*[\.\)]\s*', s)
    cleaned = [p.strip().rstrip(',').strip() for p in pieces if p.strip()]
    if not cleaned and s:
        cleaned = [p.strip() for p in s.split(',') if len(p.strip()) > 2]
    return cleaned if cleaned else [s]

applications = {}
inventors_pool = {}
prosecution_dates = {}
forms_tracker = {}
comments_queries = []

print("=== STARTING HIGH-EFFICIENCY DATA EXTRACTION ===")

# 1. LR_ALL PATENT LIST data.xlsx
lr_path = os.path.join(BASE_DIR, "LR_ALL PATENT LIST data.xlsx")
if os.path.exists(lr_path):
    print(f"Reading: {lr_path}")
    wb = openpyxl.load_workbook(lr_path, read_only=True, data_only=True)
    sheet = wb['DATA']
    
    header = None
    hmap = {}
    empty_streak = 0
    
    for row in sheet.iter_rows(values_only=True):
        if not header:
            if any(row):
                header = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(row)]
                hmap = {h: i for i, h in enumerate(header)}
            continue
        
        if not any(c is not None for c in row):
            empty_streak += 1
            if empty_streak > 15:
                break
            continue
        empty_streak = 0
        
        def get_val(col_name):
            idx = hmap.get(col_name)
            return row[idx] if idx is not None and idx < len(row) else None

        code = clean_str(get_val('Code'))
        if not code:
            continue
            
        prefix, is_mahe, category = classify_prefix(code)
        title = clean_str(get_val('Title')) or 'Untitled Patent Application'
        app_num = clean_str(get_val('Application Number'))
        app_name = clean_str(get_val('Applicant Name')) or ('Manipal Academy of Higher Education' if is_mahe else 'Individual / Corporate')
        app_addr = clean_str(get_val('Applicant address\n')) or clean_str(get_val('Applicant address'))
        status = clean_str(get_val('STATUS')) or 'PS Filed'
        color_code = clean_str(get_val('Color Code Status')) or 'GREEN'
        if color_code.upper() not in {'PURPLE', 'GREEN', 'RED', 'YELLOW', 'ORANGE', 'BLUE', 'GREY'}:
            color_code = 'GREEN'
        is_abandoned = (color_code.upper() == 'RED') or ('abandon' in status.lower())
        
        raw_inv_names = split_numbered_list(get_val('Inventor Name'))
        raw_inv_addrs = split_numbered_list(get_val('Inventor address'))
        raw_inv_insts = split_numbered_list(get_val('INVENTOR Institution'))
        raw_inv_desigs = split_numbered_list(get_val('Inventor Designation'))
        inv_email = clean_str(get_val('Inventor Email'))
        inv_phone = clean_str(get_val('Phone No.'))
        
        inv_summary_list = []
        for i, name in enumerate(raw_inv_names):
            name_clean = name.strip()
            if not name_clean:
                continue
            inv_summary_list.append(name_clean)
            key = f"{name_clean.lower()}_{inv_email.lower() if i==0 else ''}"
            if key not in inventors_pool:
                inventors_pool[key] = {
                    "id": f"inv_{len(inventors_pool)+1:05d}",
                    "full_name": name_clean,
                    "email": inv_email if i==0 else None,
                    "phone": inv_phone if i==0 else None,
                    "institution_department": raw_inv_insts[i] if i < len(raw_inv_insts) else (raw_inv_insts[0] if raw_inv_insts else None),
                    "designation": raw_inv_desigs[i] if i < len(raw_inv_desigs) else None,
                    "address": raw_inv_addrs[i] if i < len(raw_inv_addrs) else None
                }

        ps_date = parse_date(get_val('date of filing (PS)'))
        cs_date = parse_date(get_val('Date of filing (CS)'))
        post_date = parse_date(get_val('Post-Dated Date'))
        cs_deadline = parse_date(get_val('CS DEADLINE')) or add_months_date(post_date or ps_date, 12)
        
        spec_type = 'PS'
        if ps_date and cs_date:
            spec_type = 'PS_TO_CS'
        elif cs_date and not ps_date:
            spec_type = 'CS'

        applications[code] = {
            "id": f"app_{len(applications)+1:05d}",
            "project_code": code,
            "application_number": app_num,
            "patent_number": None,
            "title": title,
            "record_category": category,
            "specification_type": spec_type,
            "classification": prefix,
            "applicant_name": app_name,
            "applicant_address": app_addr,
            "is_mahe": is_mahe,
            "entity_category": "EDUCATIONAL_INSTITUTION" if is_mahe else "LARGE_ENTITY",
            "external_attorney": None,
            "assigned_to": clean_str(get_val('Assigned to ')) or clean_str(get_val('Assigned to')),
            "filed_by": clean_str(get_val('Filed BY')),
            "application_filing_receipt_status": clean_str(get_val('APPLICATION FILING RECEIPT')) or 'RECEIVED',
            "communication_status": 'OPEN',
            "current_status": status,
            "is_abandoned": is_abandoned,
            "color_code_status": color_code.upper(),
            "workdrive_folder_url": None,
            "inventors_summary": ", ".join(inv_summary_list) if inv_summary_list else None
        }

        pub_date = parse_date(get_val('Date of Publication'))
        fer_date = parse_date(get_val('Date of FER Notice'))
        fer_reply = parse_date(get_val('Date of filing FER Reply'))
        fer_deadline = parse_date(get_val('FER Reply Deadline')) or add_months_date(fer_date, 6)
        hearing_notice = parse_date(get_val('Date of Hearing notice'))
        hearing_date = parse_date(get_val('Date of hearing'))
        ws_deadline = parse_date(get_val('WS Deadline')) or add_days_date(hearing_date, 15)
        ws_date = parse_date(get_val('Date of filing WS'))
        grant_date = parse_date(get_val('Date of Grant'))
        refusal_date = parse_date(get_val('Date of refusal'))
        renewal_date = parse_date(get_val('Date of Renewal'))
        next_renewal = parse_date(get_val('Date of Next Renewal'))
        rfe_date = parse_date(get_val('Form-18/18A date of filing'))

        priority_date = ps_date or cs_date
        rfe_deadline = None
        if priority_date:
            try:
                p_dt = datetime.strptime(priority_date, '%Y-%m-%d').date()
                mos = 31 if p_dt >= RULE_CUTOFF else 48
                rfe_deadline = add_months_date(priority_date, mos)
            except Exception:
                pass

        prosecution_dates[code] = {
            "ps_filing_date": ps_date,
            "post_date_allowed": clean_str(get_val('POST DATE ALLOWED')).lower() in {'yes', 'true'},
            "post_dated_date": post_date,
            "remaining_post_date_days": 181,
            "cs_filing_date": cs_date,
            "cs_deadline": cs_deadline,
            "publication_date": pub_date,
            "publication_number": clean_str(get_val('Publication number')),
            "rfe_filed": bool(rfe_date),
            "rfe_filing_date": rfe_date,
            "rfe_deadline": rfe_deadline,
            "fer_notice_date": fer_date,
            "fer_reply_filing_date": fer_reply,
            "fer_reply_deadline": fer_deadline,
            "hearing_notice_date": hearing_notice,
            "hearing_date": hearing_date,
            "ws_filing_date": ws_date,
            "ws_deadline": ws_deadline,
            "grant_date": grant_date,
            "refusal_date": refusal_date,
            "last_renewal_date": renewal_date,
            "next_renewal_date": next_renewal
        }

        f1_deadline = parse_date(get_val('Form-1 deadline')) or add_months_date(ps_date or cs_date, 6)
        f1_date = parse_date(get_val('Form-1 date of filing'))
        f26_deadline = parse_date(get_val('Form 26 Deadline')) or add_months_date(ps_date or cs_date, 3)
        f26_date = parse_date(get_val('Form-26 date of filing'))

        forms_tracker[code] = {
            "form_1_filing_date": f1_date,
            "form_1_deadline": f1_deadline,
            "f1_delayed_by": clean_str(get_val('F1 Delayed by')),
            "f1_status": 'Delayed' if (f1_date and f1_deadline and f1_date > f1_deadline) else ('Within Deadline' if f1_date else 'Pending'),
            "form_26_filing_date": f26_date,
            "form_26_deadline": f26_deadline,
            "f26_delayed_by": clean_str(get_val('F26 Delayed by')),
            "f26_status": 'Delayed' if (f26_date and f26_deadline and f26_date > f26_deadline) else ('Within Deadline' if f26_date else 'Pending'),
            "form_3_filing_date": parse_date(get_val('FORM 3 date of filing')),
            "form_3_deadline": parse_date(get_val('FORM 3 deadline')),
            "form_9_filing_date": parse_date(get_val('Form-9 date of filing')),
            "form_18_filing_date": rfe_date,
            "form_6_filing_date": parse_date(get_val('Form 6- Date of Filing')),
            "form_13_filing_date": parse_date(get_val('Form 13 date of filing')),
            "form_13_particulars": clean_str(get_val('Form 13 Particulars')),
            "form_4_fer_filing_date": parse_date(get_val('Form 4 FER Date of filing')),
            "form_4_ws_filing_date": parse_date(get_val('Form 4 for WS Date of filing')),
            "form_4_form_3_filing_date": parse_date(get_val('Form 4 for Form 3')),
            "form_8_filing_date": parse_date(get_val('Form 8 Date of filing')),
            "form_8a_filing_date": parse_date(get_val('Form 8A Date of filling')),
            "form_24_filing_date": parse_date(get_val('Form 24 Date of filing')),
            "form_27_filing_date": parse_date(get_val('Form 27 Date of filing')),
            "petition_137_filed": bool(clean_str(get_val('Petition under 137'))),
            "petition_138_filed": bool(clean_str(get_val('Petition under 138')))
        }

print(f"Loaded {len(applications)} from LR sheet.")

# 2. Ingest Drafters_Dashbaord data.xlsx (Intake/IDF)
drafters_path = os.path.join(BASE_DIR, "Drafters_Dashbaord data.xlsx")
if os.path.exists(drafters_path):
    print(f"Reading: {drafters_path}")
    wb2 = openpyxl.load_workbook(drafters_path, read_only=True, data_only=True)
    drafter_sheets = [s for s in wb2.sheetnames if s not in ['Dashboard', 'Drafter_List', 'Performance', 'All_Purple', 'All_Yellow', 'All_Review', 'All_Orange', 'All_Green']]
    
    for dname in drafter_sheets:
        s = wb2[dname]
        d_hdr = None
        dhmap = {}
        empty_streak = 0
        for r in s.iter_rows(values_only=True):
            if not d_hdr:
                if any(r):
                    d_hdr = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(r)]
                    dhmap = {h: i for i, h in enumerate(d_hdr)}
                continue
            if not any(c is not None for c in r):
                empty_streak += 1
                if empty_streak > 15:
                    break
                continue
            empty_streak = 0
            
            def get_d(col_name):
                idx = dhmap.get(col_name)
                return r[idx] if idx is not None and idx < len(r) else None

            code = clean_str(get_d('AR CODE'))
            if not code or len(code) > 25:
                continue
            inv_name = clean_str(get_d('INVENTOR NAME & POINT OF CONTACT'))
            case_status = clean_str(get_d('CASE STATUS')) or 'PURPLE'
            
            if code not in applications:
                prefix, is_mahe, category = classify_prefix(code)
                applications[code] = {
                    "id": f"app_{len(applications)+1:05d}",
                    "project_code": code,
                    "application_number": None,
                    "patent_number": None,
                    "title": f"Invention {code} - {inv_name or 'Patent Application'}",
                    "record_category": category,
                    "specification_type": "PS",
                    "classification": prefix,
                    "applicant_name": "Manipal Academy of Higher Education" if is_mahe else "Independent",
                    "applicant_address": None,
                    "is_mahe": is_mahe,
                    "entity_category": "EDUCATIONAL_INSTITUTION" if is_mahe else "INDIVIDUAL",
                    "external_attorney": None,
                    "assigned_to": dname,
                    "filed_by": None,
                    "application_filing_receipt_status": "PENDING",
                    "communication_status": "PENDING_INVENTOR" if case_status == 'PURPLE' else 'OPEN',
                    "current_status": f"IDF - {case_status}",
                    "is_abandoned": False,
                    "color_code_status": case_status if case_status in {'PURPLE', 'YELLOW', 'ORANGE', 'GREEN'} else 'PURPLE',
                    "workdrive_folder_url": None,
                    "inventors_summary": inv_name
                }
                prosecution_dates[code] = {
                    "ps_filing_date": None,
                    "post_date_allowed": False,
                    "post_dated_date": None,
                    "remaining_post_date_days": 181,
                    "cs_filing_date": None,
                    "cs_deadline": None,
                    "publication_date": None,
                    "publication_number": None,
                    "rfe_filed": False,
                    "rfe_filing_date": None,
                    "rfe_deadline": None,
                    "fer_notice_date": None,
                    "fer_reply_filing_date": None,
                    "fer_reply_deadline": None,
                    "hearing_notice_date": None,
                    "hearing_date": None,
                    "ws_filing_date": None,
                    "ws_deadline": None,
                    "grant_date": None,
                    "refusal_date": None,
                    "last_renewal_date": None,
                    "next_renewal_date": None
                }
                forms_tracker[code] = {
                    "form_1_filing_date": None,
                    "form_1_deadline": None,
                    "f1_delayed_by": None,
                    "f1_status": "Pending",
                    "form_26_filing_date": None,
                    "form_26_deadline": None,
                    "f26_delayed_by": None,
                    "f26_status": "Pending",
                    "form_3_filing_date": None,
                    "form_3_deadline": None,
                    "form_9_filing_date": None,
                    "form_18_filing_date": None,
                    "form_6_filing_date": None,
                    "form_13_filing_date": None,
                    "form_13_particulars": None,
                    "form_4_fer_filing_date": None,
                    "form_4_ws_filing_date": None,
                    "form_4_form_3_filing_date": None,
                    "form_8_filing_date": None,
                    "form_8a_filing_date": None,
                    "form_24_filing_date": None,
                    "form_27_filing_date": None,
                    "petition_137_filed": False,
                    "petition_138_filed": False
                }

print(f"Total applications normalized: {len(applications)}")

# Save to output JSON files
with open(os.path.join(OUTPUT_DIR, "applications.json"), "w", encoding="utf-8") as f:
    json.dump(list(applications.values()), f, indent=2)

with open(os.path.join(OUTPUT_DIR, "inventors.json"), "w", encoding="utf-8") as f:
    json.dump(list(inventors_pool.values()), f, indent=2)

with open(os.path.join(OUTPUT_DIR, "prosecution_dates.json"), "w", encoding="utf-8") as f:
    json.dump(prosecution_dates, f, indent=2)

with open(os.path.join(OUTPUT_DIR, "forms_tracker.json"), "w", encoding="utf-8") as f:
    json.dump(forms_tracker, f, indent=2)

# Build unified consolidated table records for the frontend
master_records = []
for code, app in applications.items():
    dates = prosecution_dates.get(code, {})
    forms = forms_tracker.get(code, {})
    rec = {**app, **dates, **forms}
    master_records.append(rec)

# Also copy into frontend/src/data/master_records.json
fe_data_dir = os.path.join(BASE_DIR, "frontend", "src", "data")
os.makedirs(fe_data_dir, exist_ok=True)
with open(os.path.join(OUTPUT_DIR, "master_records.json"), "w", encoding="utf-8") as f:
    json.dump(master_records, f)

with open(os.path.join(fe_data_dir, "master_records.json"), "w", encoding="utf-8") as f:
    json.dump(master_records, f)

print(f"SUCCESS: {len(master_records)} master records generated and synced to frontend/src/data/master_records.json ({os.path.getsize(os.path.join(fe_data_dir, 'master_records.json'))//1024} KB).")
